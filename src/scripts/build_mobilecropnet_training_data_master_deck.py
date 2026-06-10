#!/usr/bin/env python3
"""Build the MobileCropNet training-data master academic PPTX deck."""

from __future__ import annotations

import argparse
import json
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from PIL import Image

try:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.util import Inches, Pt
except ImportError as exc:  # pragma: no cover - exercised only when dependency is missing.
    raise SystemExit(
        "Missing dependency: python-pptx. Install it with "
        "/media/jyju25/Disk_JY/Projects_26/Venvs/ImageCropping_Py310/bin/python "
        "-m pip install python-pptx"
    ) from exc


PROJECT_ROOT = Path(__file__).resolve().parents[2]
ASSET_DIR = PROJECT_ROOT / "Implement_Docs" / "assets_mobilecropnet_training_data_master_20260424"
DEFAULT_OUTPUT = ASSET_DIR / "MobileCropNet_Training_Data_Generation_Master_Report_KO_2026-04-24.pptx"
DEFAULT_VALIDATION = ASSET_DIR / "MobileCropNet_Training_Data_Generation_Master_Report_KO_2026-04-24_deck_validation.json"

SLIDE_W = 13.333333
SLIDE_H = 7.5

PAPER = "F7F8FB"
INK = "152033"
MUTED = "5F6B7A"
GRID = "DDE3EC"
WHITE = "FFFFFF"
NAVY = "1F4E79"
TEAL = "2E8C7D"
ORANGE = "C26A2E"
GREEN = "4F8A53"
PURPLE = "7357A6"
RED = "B54A4A"
GOLD = "C49A2C"

FONT_KO = "Noto Sans CJK KR"
FONT_LATIN = "Aptos"
REPORT_NAME = "MobileCropNet Training Data Generation Master Report, 2026-04-24"


def rgb(hex_color: str) -> RGBColor:
    value = hex_color.strip().lstrip("#")
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def rel(path: Path) -> str:
    try:
        return str(path.relative_to(PROJECT_ROOT))
    except ValueError:
        return str(path)


def asset(name: str) -> Path:
    return ASSET_DIR / name


@dataclass(frozen=True)
class AssetSpec:
    path: Path
    caption: str = ""


@dataclass(frozen=True)
class SlideSpec:
    section: str
    title: str
    message: str
    bullets: tuple[str, ...] = ()
    assets: tuple[AssetSpec, ...] = ()
    kind: str = "image_bullets"
    stats: tuple[tuple[str, str], ...] = ()
    flow: tuple[str, ...] = ()
    callout: str = ""
    color: str = NAVY


def core_slides() -> list[SlideSpec]:
    return [
        SlideSpec(
            section="Core",
            title="MobileCropNet 학습 데이터 생성",
            message="온디바이스 이미지 크롭핑을 위한 후보 중심, Teacher 라우팅, 정책 인식 라벨 팩토리",
            bullets=(
                "대상 문서: Implement_Docs/MobileCropNet_Training_Data_Generation_Master_Report_KO_2026-04-24.md",
                "학술보고용 core deck + appendix 구조로 렌더링",
                "핵심 결론: teacher ceiling보다 student internalization이 현재 병목",
            ),
            kind="title",
            flow=("Source", "C1-C7", "Route", "Candidates", "Teacher", "Labels", "Benchmark"),
            color=NAVY,
        ),
        SlideSpec(
            section="Core",
            title="한 장 요약",
            message="현재 문제는 좋은 teacher가 없는 것이 아니라, no-prior student가 route, subject, policy, explanation 구조를 충분히 내재화하지 못하는 것이다.",
            bullets=(
                "production_final_hybrid는 equal-4 raw 0.882696, z 1.148158로 강한 teacher ceiling을 보임",
                "최종 student shortlist 19/19가 route-collapse hard gate 실패",
                "다음 개선은 route-balanced sampling, subject-aware proposal, portrait-sensitive checklist에 집중",
            ),
            assets=(AssetSpec(asset("fig_final_gate_pass_fail_matrix.png"), "Final release gate pass/fail matrix"),),
            kind="hero_image",
            color=RED,
        ),
        SlideSpec(
            section="Core",
            title="왜 단일 BBox 회귀가 아닌가",
            message="동일 이미지라도 target AR, 인물/객체/장면 구조, 텍스트와 copy-space, 안전한 잘림 여부에 따라 정답 crop이 달라진다.",
            bullets=(
                "좌표 회귀: 빠르지만 다중 해와 action policy를 설명하기 어렵다",
                "후보 순위화: 상대 선호를 다루지만 safety/reject provenance가 약하다",
                "MOS/VLM signal: 유용하지만 product-safe label contract로 번역되어야 한다",
            ),
            kind="landscape",
            flow=("Coordinate\nRegression", "Candidate\nRanking", "MOS\nBenchmarks", "VLM / Intent", "Subject-aware\nComposition", "MobileCropNet\nLabel Factory"),
            color=TEAL,
        ),
        SlideSpec(
            section="Core",
            title="문제 정식화",
            message="학습 label은 하나의 box가 아니라 candidate availability, local ranking, policy decision, explanation consistency를 동시에 담는 contract다.",
            bullets=(
                "후보 집합 C(I,r)를 만든 뒤 image-targetAR 내부 ranking을 학습",
                "policy score는 baseline 대비 crop/minimal/keep action을 결정",
                "unsafe, fatal, contradiction state는 score보다 우선하는 gate",
            ),
            kind="error",
            flow=("Candidate\nError", "Ranking\nError", "Policy\nError", "Explanation\nError"),
            callout="E_total = E_candidate + E_ranking + E_policy + E_explain",
            color=PURPLE,
        ),
        SlideSpec(
            section="Core",
            title="End-to-End Label Factory",
            message="source corpus부터 benchmark feedback까지 이어지는 데이터 팩토리가 MobileCropNet의 행동 공간을 정의한다.",
            assets=(AssetSpec(asset("fig_pipeline_overview.png"), "Existing pipeline overview"),),
            kind="pipeline",
            flow=("Source\nCorpora", "Curation\n+ Splits", "C1-C7\nPrecompute", "Routing\nGuidance", "Candidate\nBank", "Teacher\nScoring", "Student\nContract", "Benchmark\nFeedback"),
            color=NAVY,
        ),
        SlideSpec(
            section="Core",
            title="Source Corpus Roles",
            message="SSTK는 product semantics를, GAIC는 public MOS/ranking benchmark signal을 제공한다. 두 source를 같은 의미로 합치면 안 된다.",
            bullets=(
                "SSTK: route, decision, checklist, safety, provenance 중심",
                "GAIC: dense candidate MOS와 public ranking 중심",
                "split은 image-level deterministic policy로 관리해 leakage를 막는다",
            ),
            kind="corpus",
            flow=("SSTK\nproduct semantics", "Image-level\nsplit", "GAIC\nMOS benchmark"),
            color=ORANGE,
        ),
        SlideSpec(
            section="Core",
            title="SSTK Factory Scale",
            message="10,000 curated image는 image-targetAR task와 sidecar supervision으로 확장되어 단순 coordinate dataset보다 훨씬 조밀한 학습 신호를 만든다.",
            stats=(
                ("Curated images", "10,000"),
                ("Batch rows", "48,766"),
                ("Pairwise rows", "151,049"),
                ("Listwise rows", "48,766"),
                ("Candidate slots", "341,665"),
                ("Candidate bank", "7,533,038"),
            ),
            assets=(AssetSpec(asset("fig_quant_sstk_factory_scale.png"), "SSTK factory supervision scale"),),
            kind="stats_image",
            color=GREEN,
        ),
        SlideSpec(
            section="Core",
            title="C1-C7 Perception Stack",
            message="C1-C7은 feature cache가 아니라 routing, candidate generation, scoring, label repair가 공유하는 관측 기반이다.",
            bullets=(
                "C1 caption/semantic, C2 object/segmentation, C3 face/pose/gaze",
                "C4 OCR/copy-space, C5 geometry, C6 portrait refinement, C7 saliency-support",
                "downstream label은 이 관측을 route, subject, checklist, risk target으로 번역",
            ),
            assets=(AssetSpec(asset("fig_c1_c7_perception_stack_panel.png"), "C1-C7 stack and downstream label uses"),),
            kind="image_bullets",
            color=TEAL,
        ),
        SlideSpec(
            section="Core",
            title="Subject Support Is Not Crop Box",
            message="C2 object box, C3 person prior, C7 support map은 서로 다른 supervision이며 winner crop과 동일시하면 안 된다.",
            bullets=(
                "support region은 teacher가 보존하려는 시각적 근거",
                "student는 runtime에서 support input 없이 image feature로 이 구조를 내재화해야 함",
                "정성 audit은 support/crop mismatch를 별도 실패 유형으로 본다",
            ),
            assets=(
                AssetSpec(asset("fig_c2_c3_c7_joint_overlay.png"), "C2/C3/C7 joint perception overlay"),
                AssetSpec(asset("fig_subject_support_map_region_examples.png"), "Subject support-map region examples"),
            ),
            kind="two_images",
            color=GREEN,
        ),
        SlideSpec(
            section="Core",
            title="Routing as Policy Control",
            message="subject_mode는 auxiliary label이 아니라 candidate seed, scoring weights, checklist applicability, action decision을 바꾸는 control variable이다.",
            bullets=(
                "portrait_single/group, object_single/multi, scene_general, copyspace_texture는 서로 다른 crop policy를 요구",
                "final shortlist는 route-collapse gate를 통과하지 못함",
                "route loss만 키우기보다 route별 hard negative와 subject-aware proposal coupling이 필요",
            ),
            assets=(AssetSpec(asset("fig_route_collapse_examples.png"), "Route supervision and collapse diagnostics"),),
            kind="route",
            flow=("subject_mode", "Candidate\nSeeds", "Scoring\nWeights", "Checklist\nApplicability", "Action\nDecision"),
            color=RED,
        ),
        SlideSpec(
            section="Core",
            title="Candidate Bank and Roles",
            message="좋은 crop이 후보 bank에 없으면 ranking, policy, explanation head가 복구할 수 없다.",
            bullets=(
                "baseline_full과 baseline_minimal은 action reference 역할",
                "matching_targets, candidate_pool, ignored, overflow는 학습과 audit의 역할을 분리",
                "Product-AR는 GAIC free-form보다 좁은 AR-conditioned candidate interface",
            ),
            assets=(
                AssetSpec(asset("fig_candidate_bank_by_ar.png"), "Candidate bank by target AR"),
                AssetSpec(asset("fig_matching_candidate_pool_ignored_overflow.png"), "Candidate role split"),
            ),
            kind="two_images",
            color=PURPLE,
        ),
        SlideSpec(
            section="Core",
            title="Teacher Lineage",
            message="T1, T6, UCTR-H stage3, public ensemble, production hybrid는 같은 score가 아니라 서로 다른 의미의 teacher lineage다.",
            bullets=(
                "T1: native product semantics와 checklist/safety",
                "UCTR-H stage3: public geometry와 GAIC ranking을 함께 학습한 deep ranker",
                "public ensemble: GAIC/CGS ranking strength, product safety는 별도 gate 필요",
            ),
            assets=(AssetSpec(asset("fig_uctr_stage3_deep_ranker_flow.png"), "UCTR stage3 deep ranker scoring flow"),),
            kind="image_bullets",
            color=ORANGE,
        ),
        SlideSpec(
            section="Core",
            title="Safe Pseudo-Label Conversion",
            message="raw public score는 강력하지만, explanation/safety contradiction 때문에 fatal demotion과 unsafe cap 이후에만 student supervision이 된다.",
            bullets=(
                "train raw public export contradiction rate 약 0.5665, fatal rate 약 0.1118",
                "safe conversion은 best safe positive, score cap, provenance preservation을 강제",
                "external score는 ranking field를 강화하지만 safety/reject semantics를 지울 수 없다",
            ),
            assets=(
                AssetSpec(asset("fig_teacher_score_disagreement_examples.png"), "Public score vs SSTK safety disagreement"),
                AssetSpec(asset("fig_safe_conversion_demoted_candidates.png"), "Safe conversion demotion/capping"),
            ),
            kind="two_images",
            color=RED,
        ),
        SlideSpec(
            section="Core",
            title="Student Training Contract",
            message="batch JSONL과 pairwise/listwise/checklist sidecar는 image-targetAR-local key로 결합되어 multi-head target을 만든다.",
            bullets=(
                "핵심 join key: (image_id, target_ar, candidate_id)",
                "pairwise는 두 후보가 sampled slot에 있을 때만 explicit preference로 주입",
                "global MOS regression이 아니라 task-local ordering으로 해석",
            ),
            assets=(
                AssetSpec(asset("fig_training_row_contract_example.png"), "Training row contract"),
                AssetSpec(asset("fig_batch_pairwise_listwise_join.png"), "Batch and sidecar join scale"),
            ),
            kind="two_images",
            color=NAVY,
        ),
        SlideSpec(
            section="Core",
            title="Product-AR and Decision Semantics",
            message="Product-AR는 deployment에 가까운 image-targetAR task이지만, branch materialization과 decision diversity를 분리해서 해석해야 한다.",
            bullets=(
                "SSTK public ensemble Product-AR는 full materialized line",
                "일부 GAIC/T1/UCTR local manifest는 smoke/fallback 수준",
                "all-crop decision label은 runtime action executor를 과소 지정한다",
            ),
            assets=(
                AssetSpec(asset("fig_product_ar_branch_materialization_status.png"), "Product-AR branch status"),
                AssetSpec(asset("fig_decision_distribution_sstk_vs_product_ar.png"), "SSTK decision/action distribution"),
            ),
            kind="two_images",
            color=GOLD,
        ),
        SlideSpec(
            section="Core",
            title="No-Prior Runtime Contract",
            message="teacher subject/support는 training supervision일 뿐 runtime input이 아니다. Product inference는 image + target_ar만 받는다.",
            bullets=(
                "subject_box_target은 model 내부 subject localization을 학습시키는 장치",
                "keep_full, minimal_crop, crop action은 runtime baseline/proposal executor와 정렬되어야 함",
                "crop score가 좋아도 subject/action/rationale가 불일치하면 deployment risk",
            ),
            assets=(
                AssetSpec(asset("fig_runtime_no_prior_action_executor_example.png"), "No-prior runtime executor"),
                AssetSpec(asset("fig_runtime_pred_subject_box_without_prior.png"), "No-prior subject-box supervision"),
            ),
            kind="two_images",
            color=TEAL,
        ),
        SlideSpec(
            section="Core",
            title="Teacher Ceiling",
            message="teacher benchmark ceiling은 이미 강하다. 다음 병목은 teacher discovery보다 transfer와 student internalization이다.",
            stats=(
                ("Hybrid equal4 raw", "0.882696"),
                ("Hybrid equal4 z", "1.148158"),
                ("UCTR stage3 raw", "0.852693"),
                ("Public ensemble raw", "0.798613"),
            ),
            assets=(
                AssetSpec(asset("fig_unified_equal4_teacher_cropper_leaderboard.png"), "Equal-4 teacher/cropper leaderboard"),
                AssetSpec(asset("fig_unified_public_benchmark_metric_matrix.png"), "Dataset-primary metric matrix"),
            ),
            kind="two_images_stats",
            color=NAVY,
        ),
        SlideSpec(
            section="Core",
            title="Student Shortlist Blocker",
            message="best-available student는 있으나 final shipping model은 아니다. public crop quality는 일부 보존되지만 route-collapse hard gate를 실패한다.",
            stats=(
                ("Shortlist entries", "19"),
                ("Route collapse", "19 / 19"),
                ("Final gate pass", "0"),
                ("Default winner", "balanced_288"),
                ("Portrait single acc", "0.01204"),
                ("Equal4 raw mean", "0.760155"),
            ),
            assets=(AssetSpec(asset("fig_quant_route_collapse_diagnostics.png"), "Student shortlist route-collapse diagnostic"),),
            kind="stats_image",
            color=RED,
        ),
        SlideSpec(
            section="Core",
            title="Qualitative Failure Taxonomy",
            message="deployment sign-off는 visual crop 품질뿐 아니라 route, action, subject-box, checklist, why-tag 일관성을 함께 본다.",
            bullets=(
                "미적 판단 전에 AR validity와 subject preservation을 먼저 확인",
                "route mismatch는 crop이 그럴듯해 보여도 1급 실패",
                "candidate failure와 ranking failure를 분리해서 원인을 기록",
            ),
            assets=(
                AssetSpec(asset("fig_qualitative_failure_taxonomy_panel.png"), "Failure taxonomy panel"),
                AssetSpec(asset("fig_headroom_lookroom_good_bad_examples.png"), "Headroom/lookroom examples"),
            ),
            kind="two_images",
            color=PURPLE,
        ),
        SlideSpec(
            section="Core",
            title="Conclusion and Next Experiments",
            message="핵심 기여는 이질적인 crop knowledge를 structured student contract로 번역한 것이다. 다음 과학적 단계는 route/action/safety gate를 통과하는 deployable student를 가르치는 것이다.",
            bullets=(
                "route-balanced sampling과 route별 hard negative로 collapse를 깨는 bounded rerun",
                "subject-aware proposal coupling과 pose-aware portrait prior 강화",
                "dense support-map supervision과 standardized qualitative review pack 승격",
            ),
            assets=(AssetSpec(asset("fig_deployment_leaderboard_summary.png"), "Deployment leaderboard signal summary"),),
            kind="hero_image",
            color=GREEN,
        ),
    ]


def appendix_slides() -> list[SlideSpec]:
    return [
        SlideSpec(
            section="Appendix",
            title="A1. Full Related Work Matrix",
            message="기존 crop supervision은 각자 강점이 다르며, MobileCropNet은 이를 structured contract로 결합한다.",
            bullets=(
                "좌표 회귀: 빠른 단일 crop 좌표",
                "후보 순위화 / GAIC MOS / VLM instruction / subject-aware composition은 서로 다른 supervision axis",
                "raw dataset metric을 product policy로 오해하지 않는 것이 핵심",
            ),
            assets=(AssetSpec(asset("fig_public_dataset_annotation_examples.png"), "Public crop benchmark annotation examples"),),
            kind="image_bullets",
            color=NAVY,
        ),
        SlideSpec(
            section="Appendix",
            title="A2. Full Objective Formulas",
            message="본문에는 핵심만 남기고 자세한 수식은 appendix에서 참조한다.",
            bullets=(
                "S_rank: aesthetic, subject, composition, external teacher의 weighted utility",
                "S_policy: rank utility에 area prior와 safety penalty를 결합",
                "listwise, pairwise, decision target은 모두 image-task-local target",
            ),
            kind="formula",
            flow=("S_rank", "S_policy", "y_rank", "p_i", "L_pair", "y_decision"),
            color=PURPLE,
        ),
        SlideSpec(
            section="Appendix",
            title="A3. SSTK Category Distribution",
            message="Full_10000 materialized corpus는 category-balanced curation을 목표로 한다.",
            stats=(
                ("Super-categories", "12"),
                ("Image files", "10,000"),
                ("people_single", "833"),
                ("people_multi", "833"),
                ("sports", "833"),
                ("architecture", "836"),
            ),
            kind="stats_only",
            color=GREEN,
        ),
        SlideSpec(
            section="Appendix",
            title="A4. Split Governance",
            message="pairwise/listwise와 Product-AR task 확장 때문에 split은 row가 아니라 image id와 seed 기준으로 관리해야 한다.",
            bullets=(
                "같은 image의 positive/negative가 split을 넘으면 ranking 평가가 과대평가된다",
                "Product-AR는 image-targetAR task를 여러 개 만들기 때문에 row split만으로는 부족",
                "branch 비교는 동일 split policy 위에서만 해석 가능",
            ),
            kind="flow",
            flow=("Image ID", "Deterministic\nSplit", "Batch JSONL", "Sidecars", "Product-AR\nTasks"),
            color=ORANGE,
        ),
        SlideSpec(
            section="Appendix",
            title="A5. C1-C7 Stage Details",
            message="각 C-stage는 downstream label의 일부로 직접 또는 간접 반영된다.",
            bullets=(
                "C1 semantic context, C2 object/segmentation, C3 face/pose/gaze",
                "C4 OCR/text/copy-space, C5 horizon/symmetry/layout",
                "C6 person refinement, C7 saliency/support attribution",
            ),
            assets=(AssetSpec(asset("fig_c7_subject_support_examples.png"), "C7 subject-support examples"),),
            kind="image_bullets",
            color=TEAL,
        ),
        SlideSpec(
            section="Appendix",
            title="A6. Route Policy Table",
            message="route family별 구성 정책과 collapse 증상은 core deck의 route-control slide를 보완한다.",
            bullets=(
                "portrait_single: face/eye-line, headroom/lookroom",
                "portrait_group: group envelope와 balance",
                "scene/copy-space: horizon, context, text/copy-space preservation",
            ),
            assets=(
                AssetSpec(asset("fig_route_subject_mode_gallery.png"), "Subject-mode routing examples"),
                AssetSpec(asset("fig_route_subject_mode_plain_gallery.png"), "Plain routing gallery"),
            ),
            kind="two_images",
            color=RED,
        ),
        SlideSpec(
            section="Appendix",
            title="A7. Candidate Algorithm",
            message="Candidate bank는 baseline, AR proposal, subject/support expansion, teacher proposal을 합친 뒤 staged NMS와 role assignment를 수행한다.",
            kind="flow",
            flow=("baseline_full", "baseline_minimal", "AR grid", "subject expansion", "support proposal", "teacher proposal", "NMS", "role assignment"),
            color=PURPLE,
        ),
        SlideSpec(
            section="Appendix",
            title="A8. Candidate Bank Numeric Table",
            message="FREE와 fixed-AR task는 서로 다른 후보 분포를 만들며, main batch는 전체 bank를 압축한 slot만 사용한다.",
            stats=(
                ("FREE total", "1,951,872"),
                ("4:3 total", "1,350,612"),
                ("16:9 total", "1,325,602"),
                ("1:1 total", "1,211,483"),
                ("3:4 total", "906,683"),
                ("9:16 total", "786,786"),
            ),
            assets=(AssetSpec(asset("fig_candidate_bank_by_ar.png"), "Target-AR candidate bank size"),),
            kind="stats_image",
            color=NAVY,
        ),
        SlideSpec(
            section="Appendix",
            title="A9. Label Role Tables",
            message="positive가 많다는 것은 쉬운 dataset이라는 뜻이 아니라 soft positive, baseline positive, hard negative audit이 공존한다는 뜻이다.",
            stats=(
                ("matching_targets", "179,494"),
                ("candidate_pool", "83,589"),
                ("ignored", "62,128"),
                ("overflow", "16,454"),
                ("top1", "48,766"),
                ("soft_positive", "120,996"),
            ),
            assets=(AssetSpec(asset("fig_matching_candidate_pool_ignored_overflow.png"), "Candidate role split"),),
            kind="stats_image",
            color=GREEN,
        ),
        SlideSpec(
            section="Appendix",
            title="A10. T1 vs UCTR Stage3",
            message="T1은 product-safe native utility, UCTR-H stage3는 public benchmark utility에 강한 deep crop ranker다.",
            bullets=(
                "T1: C1-C7 evidence, route, checklist, safety, decision type을 직접 보존",
                "UCTR-H: FCDB/GNMC IoU, CPC pairwise, GAIC MOS/ranking을 통합 학습",
                "UCTR branch도 safe conversion을 통과해야 product label이 된다",
            ),
            assets=(AssetSpec(asset("fig_uctr_stage3_deep_ranker_flow.png"), "T1 native score vs UCTR stage3"),),
            kind="image_bullets",
            color=ORANGE,
        ),
        SlideSpec(
            section="Appendix",
            title="A11. Pseudo-Label Semantic Layers",
            message="external teacher override 가능 범위는 geometry/rank/policy/safety/explanation/provenance 계층마다 다르다.",
            kind="layers",
            flow=("Geometry", "Rank score", "Policy decision", "Safety / reject", "Explanation", "Provenance"),
            color=RED,
        ),
        SlideSpec(
            section="Appendix",
            title="A12. GAIC Corrected vs Product-AR",
            message="corrected_v2b는 broad free-form ranking, product_ar_v1은 narrow AR-conditioned product task다.",
            stats=(
                ("corrected train", "2,509"),
                ("product-AR train", "14,915"),
                ("corrected cand/task", "~86"),
                ("product cand/task", "~6"),
                ("corrected pairwise", "240,864"),
                ("product pairwise", "251,239"),
            ),
            assets=(AssetSpec(asset("fig_quant_gaic_label_lineage_comparison.png"), "GAIC lineage candidate richness"),),
            kind="stats_image",
            color=PURPLE,
        ),
        SlideSpec(
            section="Appendix",
            title="A13. Public Explain-Safe Details",
            message="public ranker score는 raw truth로 쓰기보다 safety-aware distillation으로 사용하는 것이 안전하다.",
            stats=(
                ("Raw labels", "227,755"),
                ("Contradiction", "0.5665"),
                ("Fatal", "0.1118"),
                ("Top1 MOS", "4.0516"),
                ("Safe pairwise", "240,864"),
                ("Distill pairwise", "401,440"),
            ),
            assets=(AssetSpec(asset("fig_safe_conversion_demoted_candidates.png"), "Safe conversion effect"),),
            kind="stats_image",
            color=RED,
        ),
        SlideSpec(
            section="Appendix",
            title="A14. Full Teacher Leaderboard",
            message="equal-4 평균과 worst-dataset z를 함께 봐야 특정 dataset collapse를 놓치지 않는다.",
            assets=(
                AssetSpec(asset("fig_unified_equal4_raw_mean_bar.png"), "Equal-4 raw mean"),
                AssetSpec(asset("fig_unified_selected_methods_comparison.png"), "Selected method comparison"),
            ),
            kind="two_images",
            color=NAVY,
        ),
        SlideSpec(
            section="Appendix",
            title="A15. Subject-Box Follow-Up",
            message="SSTK subject-box v2는 subject IoU를 개선하지만 route-collapse 해결을 자동으로 보장하지 않는다.",
            stats=(
                ("SSTK plus Val IoU", "0.477671"),
                ("SSTK plus Direct IoU", "0.522159"),
                ("SSTK plus Neg acc", "0.740658"),
                ("SSTK plus Hit@0.5", "0.969483"),
            ),
            assets=(AssetSpec(asset("fig_quant_subject_box_v1_v2.png"), "Subject-box v2 quantitative comparison"),),
            kind="stats_image",
            color=TEAL,
        ),
        SlideSpec(
            section="Appendix",
            title="A16. Qualitative Gallery",
            message="정성 gallery는 성공 crop만이 아니라 subject FP/FN, support mismatch, portrait checklist failure를 포함해야 한다.",
            assets=(
                AssetSpec(asset("fig_subject_box_fp_fn_examples.png"), "Subject-box FP/FN examples"),
                AssetSpec(asset("fig_subject_coverage_good_bad_examples.png"), "Subject coverage good/bad examples"),
            ),
            kind="two_images",
            color=PURPLE,
        ),
        SlideSpec(
            section="Appendix",
            title="A17. Limits",
            message="현 상태의 한계는 명확히 보고해야 한다. 특히 Product-AR smoke/full 혼재와 route collapse를 과장 없이 분리한다.",
            bullets=(
                "Product-AR artifact 상태가 branch별로 full/smoke/fallback으로 혼재",
                "final shortlist 19/19 route-collapse hard gate 실패",
                "raw public score를 global regression truth로 쓰면 product semantics가 손상",
                "dense subject-support map supervision은 아직 충분히 활용되지 않음",
            ),
            kind="bullets",
            color=RED,
        ),
        SlideSpec(
            section="Appendix",
            title="A18. Artifact and Reproducibility Index",
            message="deck은 source markdown, figure manifest, generated PPTX, validation summary를 같은 asset directory 아래에서 추적한다.",
            bullets=(
                "source report: Implement_Docs/MobileCropNet_Training_Data_Generation_Master_Report_KO_2026-04-24.md",
                "figure manifest: Implement_Docs/assets_mobilecropnet_training_data_master_20260424/figure_manifest.json",
                "renderer: src/scripts/build_mobilecropnet_training_data_master_deck.py",
                "validation JSON: generated next to the PPTX",
            ),
            kind="bullets",
            color=NAVY,
        ),
    ]


def set_slide_background(slide, color: str = PAPER) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_textbox(slide, x: float, y: float, w: float, h: float, text: str, size: int = 18, color: str = INK, bold: bool = False, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.03)
    frame.margin_right = Inches(0.03)
    frame.margin_top = Inches(0.02)
    frame.margin_bottom = Inches(0.02)
    p = frame.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = FONT_KO
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return box


def add_header(slide, spec: SlideSpec, index: int, total: int) -> None:
    add_textbox(slide, 0.45, 0.24, 10.85, 0.43, spec.title, 24, INK, True)
    section_color = spec.color
    tag = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(11.45), Inches(0.25), Inches(1.35), Inches(0.32))
    tag.fill.solid()
    tag.fill.fore_color.rgb = rgb(section_color)
    tag.line.color.rgb = rgb(section_color)
    tf = tag.text_frame
    tf.clear()
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = spec.section
    r.font.name = FONT_LATIN
    r.font.size = Pt(10)
    r.font.bold = True
    r.font.color.rgb = rgb(WHITE)
    add_footer(slide, index, total)


def add_footer(slide, index: int, total: int) -> None:
    line = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(0.45), Inches(7.05), Inches(12.45), Inches(0.01))
    line.fill.solid()
    line.fill.fore_color.rgb = rgb(GRID)
    line.line.color.rgb = rgb(GRID)
    add_textbox(slide, 0.45, 7.12, 9.7, 0.22, REPORT_NAME, 7, MUTED, False)
    add_textbox(slide, 12.15, 7.12, 0.75, 0.22, f"{index}/{total}", 7, MUTED, False, PP_ALIGN.RIGHT)


def add_message(slide, message: str, y: float = 0.78, h: float = 0.58) -> None:
    box = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(y), Inches(12.25), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = rgb(WHITE)
    box.line.color.rgb = rgb(GRID)
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.LEFT
    r = p.add_run()
    r.text = message
    r.font.name = FONT_KO
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = rgb(INK)


def add_bullets(slide, bullets: Iterable[str], x: float, y: float, w: float, h: float, size: int = 14) -> None:
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.05)
    frame.margin_top = Inches(0.04)
    frame.margin_bottom = Inches(0.02)
    for i, bullet in enumerate(bullets):
        p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
        p.text = f"- {bullet}"
        p.level = 0
        p.font.name = FONT_KO
        p.font.size = Pt(size)
        p.font.color.rgb = rgb(INK)
        p.space_after = Pt(8)


def add_caption(slide, caption: str, x: float, y: float, w: float) -> None:
    if caption:
        add_textbox(slide, x, y, w, 0.24, caption, 8, MUTED, False, PP_ALIGN.CENTER)


def fit_image(slide, path: Path, x: float, y: float, w: float, h: float) -> bool:
    if not path.exists():
        draw_missing(slide, rel(path), x, y, w, h)
        return False
    with Image.open(path) as img:
        iw, ih = img.size
    box_ratio = w / h
    img_ratio = iw / ih
    if img_ratio >= box_ratio:
        new_w = w
        new_h = w / img_ratio
    else:
        new_h = h
        new_w = h * img_ratio
    px = x + (w - new_w) / 2
    py = y + (h - new_h) / 2
    slide.shapes.add_picture(str(path), Inches(px), Inches(py), width=Inches(new_w), height=Inches(new_h))
    return True


def draw_missing(slide, label: str, x: float, y: float, w: float, h: float) -> None:
    rect = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    rect.fill.solid()
    rect.fill.fore_color.rgb = rgb("FFF2F2")
    rect.line.color.rgb = rgb(RED)
    add_textbox(slide, x + 0.1, y + 0.1, w - 0.2, h - 0.2, f"Missing asset:\n{label}", 11, RED, True, PP_ALIGN.CENTER)


def add_card(slide, x: float, y: float, w: float, h: float, title: str, body: str = "", color: str = NAVY, size: int = 12) -> None:
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(WHITE)
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(1.25)
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = title if not body else f"{title}\n{body}"
    r.font.name = FONT_KO
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.color.rgb = rgb(INK)


def add_arrow(slide, x1: float, y1: float, x2: float, y2: float, color: str = MUTED) -> None:
    line = slide.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    line.line.color.rgb = rgb(color)
    line.line.width = Pt(1.2)
    line.line.end_arrowhead = True


def add_stats_grid(slide, stats: Iterable[tuple[str, str]], x: float, y: float, w: float, h: float, color: str) -> None:
    stats = list(stats)
    cols = 2 if len(stats) <= 6 else 3
    rows = (len(stats) + cols - 1) // cols
    gap = 0.12
    cell_w = (w - gap * (cols - 1)) / cols
    cell_h = (h - gap * (rows - 1)) / rows
    for idx, (label, value) in enumerate(stats):
        col = idx % cols
        row = idx // cols
        cx = x + col * (cell_w + gap)
        cy = y + row * (cell_h + gap)
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, Inches(cx), Inches(cy), Inches(cell_w), Inches(cell_h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(WHITE)
        shape.line.color.rgb = rgb(color)
        tf = shape.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r1 = p.add_run()
        r1.text = f"{value}\n"
        r1.font.name = FONT_LATIN
        r1.font.size = Pt(20 if len(value) <= 9 else 15)
        r1.font.bold = True
        r1.font.color.rgb = rgb(color)
        r2 = p.add_run()
        r2.text = label
        r2.font.name = FONT_KO
        r2.font.size = Pt(9)
        r2.font.bold = False
        r2.font.color.rgb = rgb(MUTED)


def render_title(slide, spec: SlideSpec) -> None:
    add_textbox(slide, 0.7, 0.52, 11.9, 0.82, spec.title, 34, INK, True, PP_ALIGN.CENTER)
    add_textbox(slide, 1.2, 1.36, 10.9, 0.55, spec.message, 18, NAVY, True, PP_ALIGN.CENTER)
    x = 0.75
    y = 2.35
    card_w = 1.55
    for idx, label in enumerate(spec.flow):
        add_card(slide, x + idx * 1.72, y, card_w, 0.72, label, "", [NAVY, TEAL, GREEN, ORANGE, PURPLE, RED, GOLD][idx % 7], 10)
        if idx < len(spec.flow) - 1:
            add_arrow(slide, x + idx * 1.72 + card_w, y + 0.36, x + (idx + 1) * 1.72, y + 0.36)
    add_bullets(slide, spec.bullets, 1.25, 4.05, 10.8, 1.35, 14)
    add_footer(slide, 1, len(all_slides()))


def render_image_bullets(slide, spec: SlideSpec) -> list[str]:
    add_message(slide, spec.message)
    missing: list[str] = []
    if spec.assets:
        ok = fit_image(slide, spec.assets[0].path, 0.65, 1.65, 7.0, 4.9)
        if not ok:
            missing.append(rel(spec.assets[0].path))
        add_caption(slide, spec.assets[0].caption, 0.65, 6.58, 7.0)
        add_bullets(slide, spec.bullets, 8.05, 1.72, 4.5, 4.8, 13)
    else:
        add_bullets(slide, spec.bullets, 1.1, 1.65, 11.0, 4.8, 15)
    return missing


def render_two_images(slide, spec: SlideSpec) -> list[str]:
    add_message(slide, spec.message)
    missing: list[str] = []
    positions = [(0.65, 1.63, 5.9, 4.85), (6.85, 1.63, 5.9, 4.85)]
    for asset_spec, (x, y, w, h) in zip(spec.assets, positions):
        if not fit_image(slide, asset_spec.path, x, y, w, h):
            missing.append(rel(asset_spec.path))
        add_caption(slide, asset_spec.caption, x, 6.56, w)
    return missing


def render_hero_image(slide, spec: SlideSpec) -> list[str]:
    add_message(slide, spec.message)
    missing: list[str] = []
    if spec.assets:
        if not fit_image(slide, spec.assets[0].path, 0.75, 1.55, 7.45, 5.0):
            missing.append(rel(spec.assets[0].path))
        add_caption(slide, spec.assets[0].caption, 0.75, 6.57, 7.45)
    add_bullets(slide, spec.bullets, 8.55, 1.62, 4.0, 4.8, 12)
    return missing


def render_stats_image(slide, spec: SlideSpec) -> list[str]:
    add_message(slide, spec.message)
    add_stats_grid(slide, spec.stats, 0.68, 1.62, 4.2, 4.8, spec.color)
    missing: list[str] = []
    if spec.assets:
        if not fit_image(slide, spec.assets[0].path, 5.2, 1.58, 7.3, 5.0):
            missing.append(rel(spec.assets[0].path))
        add_caption(slide, spec.assets[0].caption, 5.2, 6.58, 7.3)
    return missing


def render_two_images_stats(slide, spec: SlideSpec) -> list[str]:
    add_message(slide, spec.message)
    add_stats_grid(slide, spec.stats, 0.65, 1.62, 3.45, 2.3, spec.color)
    missing: list[str] = []
    positions = [(4.35, 1.48, 4.0, 4.95), (8.55, 1.48, 4.1, 4.95)]
    for asset_spec, (x, y, w, h) in zip(spec.assets, positions):
        if not fit_image(slide, asset_spec.path, x, y, w, h):
            missing.append(rel(asset_spec.path))
        add_caption(slide, asset_spec.caption, x, 6.53, w)
    return missing


def render_stats_only(slide, spec: SlideSpec) -> list[str]:
    add_message(slide, spec.message)
    add_stats_grid(slide, spec.stats, 1.0, 1.65, 11.25, 4.85, spec.color)
    return []


def render_bullets(slide, spec: SlideSpec) -> list[str]:
    add_message(slide, spec.message)
    add_bullets(slide, spec.bullets, 1.0, 1.55, 11.2, 4.9, 15)
    return []


def render_flow(slide, spec: SlideSpec) -> list[str]:
    add_message(slide, spec.message)
    labels = list(spec.flow)
    start_x = 0.7
    y = 3.05
    available = 11.9
    gap = 0.18
    card_w = min(1.65, (available - gap * (len(labels) - 1)) / len(labels))
    total_w = card_w * len(labels) + gap * (len(labels) - 1)
    x = start_x + (available - total_w) / 2
    palette = [spec.color, TEAL, GREEN, ORANGE, PURPLE, GOLD, RED, NAVY]
    for idx, label in enumerate(labels):
        add_card(slide, x + idx * (card_w + gap), y, card_w, 0.82, label, "", palette[idx % len(palette)], 9)
        if idx < len(labels) - 1:
            add_arrow(slide, x + idx * (card_w + gap) + card_w, y + 0.41, x + (idx + 1) * (card_w + gap), y + 0.41)
    return []


def render_pipeline(slide, spec: SlideSpec) -> list[str]:
    add_message(slide, spec.message)
    labels = list(spec.flow)
    x0, y0 = 0.55, 2.0
    card_w, card_h, gap = 1.45, 0.66, 0.13
    palette = [NAVY, TEAL, GREEN, ORANGE, PURPLE, RED, GOLD, NAVY]
    for idx, label in enumerate(labels):
        x = x0 + idx * (card_w + gap)
        add_card(slide, x, y0, card_w, card_h, label, "", palette[idx % len(palette)], 8)
        if idx < len(labels) - 1:
            add_arrow(slide, x + card_w, y0 + card_h / 2, x + card_w + gap, y0 + card_h / 2)
    missing: list[str] = []
    if spec.assets:
        if not fit_image(slide, spec.assets[0].path, 1.0, 3.08, 11.4, 3.18):
            missing.append(rel(spec.assets[0].path))
        add_caption(slide, spec.assets[0].caption, 1.0, 6.42, 11.4)
    return missing


def render_landscape(slide, spec: SlideSpec) -> list[str]:
    add_message(slide, spec.message)
    labels = list(spec.flow)
    center_x, center_y = 6.7, 3.95
    add_card(slide, center_x - 1.35, center_y - 0.55, 2.7, 1.1, labels[-1], "structured contract", spec.color, 13)
    positions = [(1.0, 2.0), (4.0, 1.65), (8.2, 1.65), (10.8, 2.25), (2.1, 5.0)]
    colors = [NAVY, TEAL, ORANGE, GREEN, PURPLE]
    for idx, (label, (x, y)) in enumerate(zip(labels[:-1], positions)):
        add_card(slide, x, y, 1.85, 0.78, label, "", colors[idx], 9)
        add_arrow(slide, x + 1.85, y + 0.39, center_x - 1.35, center_y)
    return []


def render_error(slide, spec: SlideSpec) -> list[str]:
    add_message(slide, spec.message)
    add_card(slide, 0.9, 2.75, 1.8, 1.0, "Image I\n+ target r", "", NAVY, 12)
    add_arrow(slide, 2.7, 3.25, 3.45, 3.25)
    add_card(slide, 3.45, 2.75, 1.8, 1.0, "Candidate set\nC(I,r)", "", TEAL, 12)
    start_x = 5.95
    colors = [GREEN, ORANGE, PURPLE, RED]
    for idx, label in enumerate(spec.flow):
        add_card(slide, start_x + idx * 1.65, 2.62, 1.35, 1.25, label, "", colors[idx], 9)
        if idx == 0:
            add_arrow(slide, 5.25, 3.25, start_x, 3.25)
    add_textbox(slide, 1.25, 5.25, 10.8, 0.5, spec.callout, 20, spec.color, True, PP_ALIGN.CENTER)
    return []


def render_corpus(slide, spec: SlideSpec) -> list[str]:
    add_message(slide, spec.message)
    add_card(slide, 0.9, 2.0, 3.2, 2.55, "SSTK", "route / decision\nchecklist / safety\nprovenance", NAVY, 16)
    add_card(slide, 9.25, 2.0, 3.2, 2.55, "GAIC", "dense candidates\nMOS / public ranking\nofficial split", ORANGE, 16)
    add_card(slide, 5.05, 2.35, 3.25, 1.85, "Image-level\ndeterministic split", "leakage guard", GREEN, 15)
    add_arrow(slide, 4.1, 3.27, 5.05, 3.27)
    add_arrow(slide, 9.25, 3.27, 8.3, 3.27)
    add_bullets(slide, spec.bullets, 1.4, 5.2, 10.8, 1.15, 12)
    return []


def render_route(slide, spec: SlideSpec) -> list[str]:
    add_message(slide, spec.message)
    add_card(slide, 0.78, 2.0, 2.2, 1.2, "subject_mode", "policy prior", spec.color, 15)
    modules = spec.flow[1:]
    for idx, label in enumerate(modules):
        x = 3.55 + idx * 2.15
        add_card(slide, x, 1.95, 1.6, 1.32, label, "", [TEAL, ORANGE, PURPLE, GREEN][idx], 9)
        add_arrow(slide, 2.98, 2.6, x, 2.6)
    add_card(slide, 3.6, 4.62, 5.95, 0.78, "Failure path: collapse to object_single", "wrong candidates + weak portrait checklist + wrong action", RED, 12)
    add_arrow(slide, 1.88, 3.2, 5.0, 4.62, RED)
    missing: list[str] = []
    if spec.assets:
        if not fit_image(slide, spec.assets[0].path, 9.8, 4.15, 2.9, 1.95):
            missing.append(rel(spec.assets[0].path))
    return missing


def render_formula(slide, spec: SlideSpec) -> list[str]:
    add_message(slide, spec.message)
    formulas = [
        ("S_rank", "aesthetic + subject + composition + external teacher"),
        ("S_policy", "S_rank + area prior - safety penalty"),
        ("y_rank", "image-task-local rank target"),
        ("p_i", "softmax listwise distribution"),
        ("L_pair", "positive-vs-negative preference"),
        ("y_decision", "keep / minimal / crop action target"),
    ]
    x0, y0 = 0.95, 1.8
    for idx, (name, body) in enumerate(formulas):
        x = x0 + (idx % 3) * 4.05
        y = y0 + (idx // 3) * 1.75
        add_card(slide, x, y, 3.55, 1.1, name, body, [NAVY, TEAL, ORANGE, GREEN, PURPLE, RED][idx], 11)
    return []


def render_layers(slide, spec: SlideSpec) -> list[str]:
    add_message(slide, spec.message)
    y = 1.65
    colors = [NAVY, TEAL, GREEN, ORANGE, RED, PURPLE]
    for idx, label in enumerate(spec.flow):
        add_card(slide, 2.55, y + idx * 0.74, 8.3, 0.55, label, "", colors[idx], 12)
    add_textbox(slide, 1.0, 6.3, 11.2, 0.35, "external teacher는 rank를 강화할 수 있지만 safety/reject와 provenance를 지울 수 없다.", 13, RED, True, PP_ALIGN.CENTER)
    return []


def render_slide(prs: Presentation, spec: SlideSpec, index: int, total: int) -> list[str]:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_background(slide)
    if spec.kind == "title":
        render_title(slide, spec)
        return []
    add_header(slide, spec, index, total)
    if spec.kind == "image_bullets":
        return render_image_bullets(slide, spec)
    if spec.kind == "two_images":
        return render_two_images(slide, spec)
    if spec.kind == "hero_image":
        return render_hero_image(slide, spec)
    if spec.kind == "stats_image":
        return render_stats_image(slide, spec)
    if spec.kind == "two_images_stats":
        return render_two_images_stats(slide, spec)
    if spec.kind == "stats_only":
        return render_stats_only(slide, spec)
    if spec.kind == "bullets":
        return render_bullets(slide, spec)
    if spec.kind == "flow":
        return render_flow(slide, spec)
    if spec.kind == "pipeline":
        return render_pipeline(slide, spec)
    if spec.kind == "landscape":
        return render_landscape(slide, spec)
    if spec.kind == "error":
        return render_error(slide, spec)
    if spec.kind == "corpus":
        return render_corpus(slide, spec)
    if spec.kind == "route":
        return render_route(slide, spec)
    if spec.kind == "formula":
        return render_formula(slide, spec)
    if spec.kind == "layers":
        return render_layers(slide, spec)
    return render_bullets(slide, spec)


def all_slides() -> list[SlideSpec]:
    return core_slides() + appendix_slides()


def build_deck(output: Path, validation: Path) -> dict[str, object]:
    specs = all_slides()
    output.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(SLIDE_W)
    prs.slide_height = Inches(SLIDE_H)
    missing_assets: list[str] = []
    for idx, spec in enumerate(specs, start=1):
        missing_assets.extend(render_slide(prs, spec, idx, len(specs)))
    prs.save(output)

    payload: dict[str, object] = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "output_pptx": rel(output),
        "validation_json": rel(validation),
        "source_report": "Implement_Docs/MobileCropNet_Training_Data_Generation_Master_Report_KO_2026-04-24.md",
        "slide_count": len(specs),
        "core_slide_count": len(core_slides()),
        "appendix_slide_count": len(appendix_slides()),
        "missing_asset_count": len(missing_assets),
        "missing_assets": sorted(set(missing_assets)),
        "sections": {
            "core": [slide.title for slide in core_slides()],
            "appendix": [slide.title for slide in appendix_slides()],
        },
    }
    validation.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return payload


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--validation", type=Path, default=DEFAULT_VALIDATION)
    return parser


def main() -> int:
    args = build_parser().parse_args()
    payload = build_deck(args.output, args.validation)
    print(json.dumps(payload, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
